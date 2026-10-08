"""Independent re-derivation of one flagged item per test.

Does NOT import common.py, clean.py or tests.py. Re-reads the source workbooks (openpyxl) and
data/contracts.csv directly, with its own header search, name normalisation and date parsing.
Writes "self_check" into outputs/analytics.json; prints one line per check; exits 1 on any failure.
"""
import csv
import difflib  # noqa: F401  (not used: candidates are never matches, so nothing to recompute)
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SPEND = ROOT / "data" / "spend"
CONTRACTS = ROOT / "data" / "contracts.csv"
OUT = ROOT / "outputs"
FULLD = OUT / "analytics-full"
ANALYTICS = OUT / "analytics.json"
COLS = ["Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name"]
BANNED = ["fraud", "irregular", "non-compliant", "noncompliant", "breach", "misconduct", "wrongdoing", "suspicious"]

checks = []


def check(test, item, ok, detail):
    checks.append({"test": test, "item": str(item), "ok": bool(ok), "detail": detail})
    print(f"[{'OK ' if ok else 'FAIL'}] {test} | {item} | {detail}")


# ---- own normalisation ------------------------------------------------------
def norm(name):
    if name is None:
        return ""
    s = str(name).upper().replace("&", " AND ")
    s = re.sub(r"['‘’ʼ`]", "", s)
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    toks = [t for t in s.split() if t not in ("LTD", "LIMITED", "PLC", "LLP")]
    if len(toks) >= 2 and toks[-1] == "ONLY" and toks[-2] in ("CHAPS", "BACS"):
        toks = toks[:-2]
    out, run = [], ""
    for t in toks:
        if len(t) == 1 and t.isalpha():
            run += t
            continue
        if run:
            out.append(run)
            run = ""
        out.append(t)
    if run:
        out.append(run)
    return " ".join(out)


def alias_set(n):
    out = {n}
    p = f" {n} "
    for sep in (" T A ", " TRADING AS ", " TA "):
        if sep in p:
            a, b = p.split(sep, 1)
            out |= {a.strip(), b.strip()}
    out.discard("")
    return out


def pence(x):
    return int(round(float(x) * 100))


# ---- own source readers -----------------------------------------------------
def load_workbooks():
    books = {}
    for f in sorted(SPEND.glob("*.xlsx")):
        if f.name.startswith("~$"):
            continue
        ws = openpyxl.load_workbook(f, data_only=True)["Data to publish"]
        rows = list(ws.iter_rows(min_row=1, values_only=True))
        hdr = None
        for i, r in enumerate(rows[:10]):
            vals = {str(v).strip() for v in r if v is not None}
            if all(c in vals for c in COLS):
                hdr = i
                break
        if hdr is None:
            check("inputs", f.name, False, "header row not found")
            continue
        header = [str(v).strip() if v is not None else "" for v in rows[hdr]]
        pos = {c: header.index(c) for c in COLS}
        data = {}
        for k, r in enumerate(rows[hdr + 1:]):
            excel_row = hdr + 1 + 1 + k
            rec = {c: (r[pos[c]] if pos[c] < len(r) else None) for c in COLS}
            if all(v is None or str(v).strip() == "" for v in rec.values()):
                continue
            d = rec["Date"]
            if isinstance(d, datetime):
                rec["day"] = d.date()
            elif isinstance(d, date):
                rec["day"] = d
            else:
                rec["day"] = datetime.strptime(str(d).strip(), "%m/%d/%y").date()
            rec["norm"] = norm(rec["Supplier name"])
            data[excel_row] = rec
        books[f.stem] = data
    return books


def load_notices():
    with open(CONTRACTS, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    scol = "Supplier [Name|Address|Ref type|Ref Number|Is SME|Is VCSE]"
    out = []
    for r in rows:
        if r["Organisation Name"].strip().lower() not in ("west berkshire council", "west berkshire"):
            continue
        packed = r[scol].strip()
        if not packed:
            continue
        for blk in packed.strip("[]").split("]["):
            name = blk.split("|")[0].strip()
            d = lambda s: datetime.strptime(s.strip(), "%d/%m/%Y").date() if s.strip() else None
            out.append({"notice_id": r["Notice Identifier"].strip(), "norm": norm(name),
                        "start": d(r["Contract start date"]), "end": d(r["Contract end date"]),
                        "awarded": float(r["Awarded Value"]) if r["Awarded Value"].strip() else 0.0})
    return out


def read_ids(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return {r["row_id"] for r in csv.DictReader(fh)}


def main():
    a = json.load(open(ANALYTICS, encoding="utf-8"))
    books = load_workbooks()
    kept_ids = read_ids(FULLD / "spend_clean.csv")
    excl_ids = read_ids(FULLD / "exclusions.csv")
    window_start = date.fromisoformat(a["window"]["first_payment"])
    months = len(books)
    factor = 12 / months
    check("window", "months", months == a["window"]["months"] and abs(factor - a["window"]["annualisation_factor"]) < 1e-9,
          f"workbooks={months}, factor={factor:g}")

    # raw rules for boundary provenance
    rules = json.load(open(OUT / "rules.json", encoding="utf-8"))
    rule_vals = set()
    for r in rules["rules"]:
        if str(r.get("clause", "")).startswith(("App A row", "App B row")):
            for k in ("threshold_low", "threshold_high"):
                if isinstance(r.get(k), (int, float)):
                    rule_vals.add(r[k])

    def fetch(test, rid, expect=None):
        stem, row = rid.rsplit(":", 1)
        rec = books.get(stem, {}).get(int(row))
        if rec is None:
            check(test, rid, False, "row_id not found in source workbook")
            return None
        if expect:
            ok = (rec["norm"] == expect.get("norm", rec["norm"])
                  and (expect.get("date") is None or str(rec["day"]) == expect["date"])
                  and (expect.get("amount") is None or pence(rec["Net amount"]) == pence(expect["amount"])))
            check(test, rid, ok, f"source supplier/date/amount = {rec['norm'][:30]!r}/{rec['day']}/{rec['Net amount']}")
        return rec

    def supplier_rows(sn):
        return [f"{st}:{er}" for st, d in books.items() for er, rec in d.items()
                if rec["norm"] == sn and f"{st}:{er}" not in excl_ids]

    tests = a["tests"]
    notices = load_notices() if CONTRACTS.exists() else []
    bounds = {b["value"]: b for b in a["rules_used"]["boundaries"]}

    for name, t in tests.items():
        if t["status"] != "run" or not t["top20"]:
            check(name, "-", True, f"status={t['status']}, nothing flagged to re-derive")
            continue
        it = t["top20"][0]
        if name == "threshold_clustering":
            T = it["boundary"]
            b = bounds[T]
            rec = fetch(name, it["row_ids"][0], {"norm": it["supplier_norm"], "date": it["date"], "amount": it["Net amount"]})
            amt = pence(rec["Net amount"])
            inband = amt >= pence(0.9 * T) and (amt < pence(T) if b["band_starts_inclusive"] else amt <= pence(T))
            check(name, f"boundary {T}", inband and T in rule_vals,
                  f"amount {amt / 100} in band below {T} (T traced to rules.json App A/B: {T in rule_vals})")
        elif name == "split_purchases":
            T = it["boundary"]
            b = bounds[T]
            recs = [fetch(name, r, {"norm": it["supplier_norm"]}) for r in it["row_ids"]]
            tot = sum(pence(r["Net amount"]) for r in recs)
            days = sorted(r["day"] for r in recs)
            span = (days[-1] - days[0]).days
            each_under = all((pence(r["Net amount"]) < pence(T)) if b["band_starts_inclusive"]
                             else (pence(r["Net amount"]) <= pence(T)) for r in recs)
            reach = tot >= pence(T) if b["band_starts_inclusive"] else tot > pence(T)
            ok = (tot == pence(it["sum_gbp"]) and span == it["span_days"] and span <= 29 and each_under
                  and reach and len(recs) == it["n"] >= 2)
            check(name, f"{it['supplier_norm'][:30]} @ {T}", ok,
                  f"sum {tot / 100} vs {it['sum_gbp']}, span {span}d, each under T {each_under}, reaches {reach}")
        elif name == "duplicates":
            recs = [fetch(name, r, {"norm": it["supplier_norm"], "amount": it["amount_gbp"]}) for r in it["row_ids"]]
            days = sorted(r["day"] for r in recs)
            gaps = [(days[i + 1] - days[i]).days for i in range(len(days) - 1)]
            ok = (len({pence(r["Net amount"]) for r in recs}) == 1 and max(gaps) <= 7 and len(recs) == it["n"]
                  and pence(it["repeat_gbp"]) == pence(it["amount_gbp"]) * (len(recs) - 1))
            check(name, it["supplier_norm"][:30], ok, f"{len(recs)} payments of {it['amount_gbp']}, max gap {max(gaps)}d")
        else:
            sn = it["supplier_norm"]
            rids = it["row_ids"]
            recs = [fetch(name, r, {"norm": sn}) for r in rids]
            tot = sum(pence(r["Net amount"]) for r in recs)
            indep = supplier_rows(sn)
            check(name, f"{sn[:30]} rows", sorted(indep) == sorted(rids),
                  f"independent kept rows for supplier {len(indep)} vs listed {len(rids)}")
            ann_ok = tot == pence(it["window_gbp"]) and abs(tot * factor / 100 - it["annualised_gbp"]) < 0.005
            check(name, f"{sn[:30]} total", ann_ok,
                  f"window {tot / 100} vs {it['window_gbp']}; annualised {tot * factor / 100:.2f} vs {it['annualised_gbp']}")
            al = alias_set(sn)
            exact = [n for n in notices if n["norm"] == sn or (al & alias_set(n["norm"]))]
            if name == "high_value_suppliers":
                q = a["rules_used"]["quote_threshold"]
                reach = (tot * factor >= q["value"] * 100) if q["inclusive"] else (tot * factor > q["value"] * 100)
                check(name, f"{sn[:30]} threshold", reach, f"annualised reaches quote_threshold {q['value']}")
            elif name == "off_contract":
                check(name, f"{sn[:30]} no exact match", not exact, f"exact notice matches found: {len(exact)}")
            elif name == "expired_notice_spend":
                ended = bool(exact) and all(n["end"] is not None and n["end"] < window_start for n in exact)
                latest = max((n["end"] for n in exact if n["end"]), default=None)
                check(name, f"{sn[:30]} expired", ended and str(latest) == str(it["latest_end"]),
                      f"{len(exact)} exact notice rows, all ended before {window_start}: {ended}; latest end {latest}")
            elif name == "spend_vs_award":
                n = [x for x in notices if x["notice_id"] == it["notice_id"]]
                n0 = n[0]
                yrs = (n0["end"] - n0["start"]).days / 365.25
                ratio = (tot * factor / 100) / (n0["awarded"] / yrs)
                ok = (len({x["norm"] for x in n}) == 1 and n0["awarded"] > 0 and round(ratio, 4) == it["ratio"]
                      and ratio >= 1.25 and any(x["notice_id"] == it["notice_id"] for x in exact))
                check(name, f"{sn[:30]} ratio", ok, f"ratio {ratio:.4f} vs {it['ratio']}; single supplier "
                                                   f"{len({x['norm'] for x in n}) == 1}; award {n0['awarded']}")

    # every flagged row_id kept, never excluded
    bad_missing, bad_excl, total = 0, 0, 0
    for name, t in tests.items():
        fl = t.get("full_list")
        if t["status"] != "run" or not fl:
            continue
        with open(ROOT / fl, encoding="utf-8", newline="") as fh:
            for r in csv.DictReader(fh):
                cell = r.get("row_ids") or r.get("row_id") or ""
                for rid in [x for x in cell.split(";") if x]:
                    total += 1
                    bad_missing += rid not in kept_ids
                    bad_excl += rid in excl_ids
    check("row_ids", "all full lists", bad_missing == 0 and bad_excl == 0,
          f"{total} row_id references; not in spend_clean: {bad_missing}; in exclusions: {bad_excl}")

    # kept + excluded = loaded, per file
    for stem, d in books.items():
        k = sum(1 for i in kept_ids if i.rsplit(":", 1)[0] == stem)
        e = sum(1 for i in excl_ids if i.rsplit(":", 1)[0] == stem)
        check("reconcile", stem, k + e == len(d) and not (kept_ids & excl_ids),
              f"kept {k} + excluded {e} = {k + e} vs loaded {len(d)}")

    # banned words
    a.pop("self_check", None)
    text = json.dumps(a).lower()
    hits = [w for w in BANNED if w in text]
    check("language", "analytics.json", not hits, f"banned words found: {hits}" if hits else "no banned words")

    passed = all(c["ok"] for c in checks)
    a["self_check"] = {"passed": passed, "checks": checks}
    with open(ANALYTICS, "w", encoding="utf-8") as fh:
        json.dump(a, fh, indent=1, default=str)
    print(f"self_check passed={passed} ({sum(c['ok'] for c in checks)}/{len(checks)})")
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main())
