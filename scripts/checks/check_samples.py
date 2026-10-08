#!/usr/bin/env python3
"""QA sample check: every sampled transaction exists in the cleaned spend data and in the council's workbook.

1. Every row of outputs/samples/T-*.csv: its row_id is in
   analytics-full/spend_clean.csv with the same month file, sheet row,
   supplier and net amount; and the council's workbook
   data/spend/<month_file>.xlsx, sheet "Data to publish", holds that supplier
   and amount at that sheet row (below the header row, located by its column
   names, never assumed).
2. Every row ID (<month_file>:<sheet row>) written in audit-program.md passes
   the same two checks against spend_clean.csv and the workbook.
3. audit-program.md and the sample files agree: every test section has a
   sample file and every sample file a test section, and each section's
   items table lists the same items with the same transaction counts as the
   file.

Usage: check_samples.py [--out DIR] [--inputs DIR] [--data DIR]
Exit 0 pass, 1 failures, 2 could not run.
"""
import importlib.util
import re
import sys
from collections import Counter
from pathlib import Path

_s = importlib.util.spec_from_file_location("_qa", Path(__file__).resolve().parent / "_qa.py")
qa = importlib.util.module_from_spec(_s)
_s.loader.exec_module(qa)

SHEET = "Data to publish"
REQUIRED_COLS = ("Service", "Expenditure category", "Narrative", "Date", "Net amount", "Supplier name")
SAMPLE_COLS = ("row_id", "month_file", "excel_row", "Supplier name", "Net amount")


def text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def amount(v):
    try:
        return float(str(v).replace(",", "").replace("£", "").strip())
    except ValueError:
        return None


class Workbooks:
    """The council's monthly workbooks, opened on first use; cells addressed by sheet row."""

    def __init__(self, folder):
        self.folder, self.books = Path(folder), {}

    def get(self, stem):
        if stem not in self.books:
            import openpyxl
            p = qa.need(self.folder / f"{stem}.xlsx", f"workbook {stem}.xlsx")
            ws = openpyxl.load_workbook(p, data_only=True)[SHEET]
            header = None
            for r in range(1, 11):
                vals = {text(ws.cell(r, k).value): k for k in range(1, ws.max_column + 1)}
                if all(col in vals for col in REQUIRED_COLS):
                    header = (r, vals["Supplier name"], vals["Net amount"])
                    break
            if header is None:
                raise qa.Stop(f"{p.name}: no header row with {', '.join(REQUIRED_COLS)} in the first 10 rows")
            self.books[stem] = (ws, header)
        return self.books[stem]


def verify(c, where, rid, spend, books, supplier=None, amt=None, month=None, row=None):
    """One transaction against spend_clean.csv and the workbook."""
    s = spend.get(rid)
    if not c.ok(s is not None, f"{where}: {rid} is not in spend_clean.csv"):
        return
    if supplier is not None:
        c.ok(text(supplier) == text(s["Supplier name"]),
             f"{where}: {rid} supplier '{supplier}' but spend_clean.csv has '{s['Supplier name']}'")
    if amt is not None:
        a, b = amount(amt), amount(s["Net amount"])
        c.ok(a is not None and b is not None and abs(a - b) <= 0.005,
             f"{where}: {rid} amount {amt} but spend_clean.csv has {s['Net amount']}")
    if month is not None:
        c.ok(month == s["month_file"] and str(row) == s["excel_row"],
             f"{where}: {rid} at {month} row {row} but spend_clean.csv has {s['month_file']} row {s['excel_row']}")
    ws, (hrow, scol, acol) = books.get(s["month_file"])
    r = int(s["excel_row"])
    if not c.ok(r > hrow, f"{where}: {rid} sheet row {r} is not below the header row {hrow}"):
        return
    wsup, wamt = ws.cell(r, scol).value, amount(ws.cell(r, acol).value)
    c.ok(text(wsup) == text(s["Supplier name"]),
         f"{where}: {rid} workbook {s['month_file']}.xlsx row {r} supplier '{text(wsup)}', "
         f"spend_clean.csv has '{s['Supplier name']}'")
    c.ok(wamt is not None and abs(wamt - float(s["Net amount"])) <= 0.005,
         f"{where}: {rid} workbook {s['month_file']}.xlsx row {r} amount {ws.cell(r, acol).value}, "
         f"spend_clean.csv has {s['Net amount']}")


def program_items(program):
    """{test_id: Counter(item number -> transactions)} from the items tables in audit-program.md."""
    out, tid = {}, None
    lines = program.splitlines()
    for i, line in enumerate(lines):
        m = re.match(r"^### (T-\d+) · ", line)
        if m:
            tid = m.group(1)
            out[tid] = Counter()
        elif line.startswith("## "):
            tid = None
        elif tid and line.startswith("| # |"):
            for row in qa.md_table(lines[i:], "#"):
                out[tid][row[0]] = int(row[-1]) if row[-1].isdigit() else -1
    return out


def body(c):
    a = qa.parser(__doc__).parse_args()
    out, inputs, data = Path(a.out), Path(a.inputs), Path(a.data)
    spend = {r["row_id"]: r for r in qa.read_csv(qa.need(inputs / "analytics-full" / qa.SPEND_FILE))}
    books = Workbooks(qa.need(data / "spend", "data/spend/"))

    # 1. sample files
    files = sorted((out / "samples").glob("T-*.csv"))
    c.ok(bool(files), "outputs/samples/ holds no T-*.csv sample files: the pack has not been built")
    sampled = {}
    for p in files:
        rows = qa.read_csv(p)
        missing = [k for k in SAMPLE_COLS if rows and k not in rows[0]]
        if not c.ok(rows and not missing, f"samples/{p.name}: " + (f"missing columns {missing}" if rows else "empty")):
            continue
        sampled[p.stem] = Counter(r.get("item", "") for r in rows)
        for n, r in enumerate(rows, start=2):
            verify(c, f"samples/{p.name}:{n}", r["row_id"], spend, books, r["Supplier name"], r["Net amount"],
                   r["month_file"], r["excel_row"])

    # 2. row IDs written in audit-program.md
    p = out / "audit-program.md"
    if not c.ok(p.exists(), "audit-program.md not found: the pack has not been built"):
        return
    program = p.read_text(encoding="utf-8")
    stems = sorted({s["month_file"] for s in spend.values()}, key=len, reverse=True)
    ids = re.compile("(?:" + "|".join(map(re.escape, stems)) + r"):\d+") if stems else None
    n_ids = 0
    for ln, line in enumerate(program.splitlines(), start=1):
        for m in ids.finditer(line) if ids else []:
            n_ids += 1
            verify(c, f"audit-program.md:{ln}", m.group(0), spend, books)
    c.note(f"{sum(sum(v.values()) for v in sampled.values())} sampled transactions in {len(files)} file(s); "
           f"{n_ids} row ID(s) written in audit-program.md")

    # 3. program items tables against the sample files
    items = program_items(program)
    for tid in sorted(set(items) | set(sampled)):
        if not c.ok(tid in sampled, f"audit-program.md test {tid} has no samples/{tid}.csv"):
            continue
        if not c.ok(tid in items, f"samples/{tid}.csv has no test section in audit-program.md"):
            continue
        c.ok(items[tid] == sampled[tid],
             f"{tid}: audit-program.md lists items {dict(items[tid])} (item: transactions) but samples/{tid}.csv "
             f"holds {dict(sampled[tid])}")


if __name__ == "__main__":
    sys.exit(qa.run("samples", body))
